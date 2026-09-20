"""CSV and XLSX writers.

CSV is the plain, greppable artifact: UTF-8, one header row, one row per ECG
sample, no BOM (a course project's data gets opened by scripts as often as by
Excel, and ``utf-8`` is what ``docs`` promise).

XLSX has to satisfy the course requirement that the waveform is viewable *in
Excel* ("可在 Excel 中进行波形显示"), which a flat table does not.  So the
workbook is three sheets:

===========  ==============================================================
``Data``     every recorded row, nothing thrown away
``Summary``  duration, counts, heart-rate and temperature statistics, drop
             and CRC counts, and the embedded chart
``ECG Series`` the decimated columns the chart actually plots
===========  ==============================================================

The chart never points at ``Data``.  Excel redraws a line chart from every
referenced point, so a ten-minute recording (600 000 samples) would lock the
spreadsheet; :func:`decimate_for_chart` reduces the plotted series to a few
thousand *extremum-preserving* points -- two per bin, the largest positive and
largest negative deviation -- so the QRS spikes survive decimation instead of
being averaged into a flat band, while ``Data`` keeps the full record.

Everything is written through openpyxl's write-only mode: a 100 000-row sheet
would otherwise cost hundreds of megabytes of live cell objects.
"""

from __future__ import annotations

import datetime as dt
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, Sequence

import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.chart import LineChart, Reference
from openpyxl.chart.series import SeriesLabel
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import i18n
from . import protocol
from .recorder import ROW_COLUMNS, RecordingSession, RecordingSummary

__all__ = [
    "CHART_SHEET",
    "DATA_SHEET",
    "SUMMARY_SHEET",
    "ExportResult",
    "decimate_for_chart",
    "export_csv",
    "export_xlsx",
    "default_stem",
    "recommended_chart_points",
    "readback_workbook",
]

DATA_SHEET = "Data"
SUMMARY_SHEET = "Summary"
CHART_SHEET = "ECG Series"

#: Keep Excel responsive: a few thousand plotted points still shows a QRS.
DEFAULT_CHART_POINTS = 4000

_HEADER_FILL = PatternFill("solid", fgColor="1F3864")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_LABEL_FONT = Font(bold=True)
_WARNING_FONT = Font(color="9C0006", bold=True)
_WARNING_FILL = PatternFill("solid", fgColor="FFC7CE")


@dataclass(frozen=True)
class ExportResult:
    path: Path
    rows: int
    sheets: tuple[str, ...] = ()
    chart_points: int = 0
    chart_series_cells: int = 0

    def describe(self) -> str:
        if self.sheets:
            return i18n.t("%s: %d rows, sheets %s, %d chart points") % (
                self.path.name,
                self.rows,
                ", ".join(i18n.t(name) for name in self.sheets),
                self.chart_points,
            )
        return i18n.t("%s: %d rows") % (self.path.name, self.rows)


def default_stem(session: RecordingSession, when: dt.datetime | None = None) -> str:
    """A filename stem that says at a glance whether it is real data.

    Deliberately not translated: a stem is a name, not interface copy, and the
    ``ecg_demo`` / ``ecg_device`` prefix is the provenance marker a script greps
    for.  ``--lang zh`` must not make a synthetic file *look* like a recording.
    """
    moment = when or dt.datetime.now()
    prefix = "ecg_demo" if session.is_demo else "ecg_device"
    return "%s_%s" % (prefix, moment.strftime("%Y%m%d-%H%M%S"))


def recommended_chart_points(row_count: int) -> int:
    """Chart resolution: the cap, but never denser than the data itself."""
    return int(max(64, min(DEFAULT_CHART_POINTS, max(1, row_count))))


def decimate_for_chart(
    times: np.ndarray, values: np.ndarray, max_points: int = DEFAULT_CHART_POINTS
) -> tuple[np.ndarray, np.ndarray, int]:
    """Extremum-preserving downsample: ``(times, values, samples_per_point)``.

    Two points per bin -- the bin's maximum and minimum -- returned in index
    order.  Peak-preserving matters for an ECG: mean-decimating a 1 kHz record by
    a factor of 15 would flatten an 80 ms QRS into a shrug, which would be worse
    than useless in a teaching workbook.
    """
    times = np.asarray(times)
    values = np.asarray(values)
    total = int(values.size)
    if total == 0:
        return times, values, 0
    if max_points <= 0 or total <= max_points:
        return times, values, 1
    bins = max(1, int(max_points) // 2)
    edges = np.linspace(0, total, bins + 1).astype(np.int64)
    keep: list[int] = []
    for start, end in zip(edges[:-1], edges[1:]):
        if end <= start:
            continue
        window = values[start:end]
        keep.append(int(start + np.argmax(window)))
        keep.append(int(start + np.argmin(window)))
    keep = sorted(set(keep))
    index = np.asarray(keep, dtype=np.int64)
    return times[index], values[index], max(1, total // len(index)) if len(index) else 1


# ------------------------------------------------------------------------ CSV
#: The one delimiter, spelled once so the header and the rows cannot disagree.
CSV_DELIMITER = ","


def csv_header_line(delimiter: str = CSV_DELIMITER) -> str:
    """The CSV header row, translated at write time.

    Each cell is the Chinese label followed by the field name it stands for.
    The label is what a reader needs; the field name is what
    ``docs/PROTOCOL.md``, ``pc_monitor/tests`` and every script that has ever
    parsed one of these files use as the join key, so it survives in both
    languages.  Column *order* comes from :data:`ROW_COLUMNS` and never changes.
    """
    return delimiter.join(i18n.t(name) for name in ROW_COLUMNS)


def export_csv(session: RecordingSession, path: str | Path) -> ExportResult:
    """Write the complete recording as UTF-8 CSV with a header row."""
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    rows = 0
    # newline="" keeps Python in charge of the line terminator, so the file is
    # byte-identical on every platform instead of CSV-CRLF-by-accident.
    with target.open("w", encoding="utf-8", newline="") as handle:
        handle.write(csv_header_line(CSV_DELIMITER))
        handle.write("\n")
        lines = session.iter_csv_lines(delimiter=CSV_DELIMITER)
        next(lines, None)  # the recorder's row is the untranslated key list
        for line in lines:
            handle.write(line)
            handle.write("\n")
            rows += 1
    return ExportResult(path=target, rows=rows)


# ----------------------------------------------------------------------- XLSX
def _summary_pairs(summary: RecordingSummary, session: RecordingSession) -> list[tuple[str, object]]:
    """The Summary sheet's key/value block, in reading order.

    Only host copy is translated here: the left-hand label and the phrases this
    tool writes itself.  Values that came off the wire -- the ``origin`` marker,
    the device identity, a ``TempState`` name, a hex flag word, an ISO
    timestamp, a version number -- are written verbatim, because the Summary has
    to agree cell for cell with the Data sheet it describes.
    """
    duration = summary.duration_s
    hrs, remainder = divmod(duration, 3600)
    minutes, seconds = divmod(int(remainder), 60)
    clock = "%02d:%02d:%04.1f" % (int(hrs), minutes, seconds + (duration - int(duration)))
    caps_text = _capabilities(session)
    return [
        (i18n.t("Recording origin"), summary.origin),
        (i18n.t("Synthetic data?"), i18n.t("YES -- DEMO MODE") if session.is_demo else i18n.t("no")),
        (i18n.t("Source"), summary.source_label or i18n.t("n/a")),
        (i18n.t("Device identity (HELLO)"), summary.device_identity),
        (i18n.t("Device capabilities"), caps_text),
        (i18n.t("Started (PC clock)"), _iso(summary.started_at)),
        (i18n.t("Stopped (PC clock)"), _iso(summary.stopped_at)),
        (i18n.t("Duration"), "%.3f s" % duration),
        (i18n.t("Duration (h:mm:ss.s)"), clock),
        (i18n.t("Sample rate (from ECG_BATCH)"), "%.0f Hz" % summary.sample_rate_hz),
        (i18n.t("Rows in Data sheet"), summary.row_count),
        (i18n.t("ECG samples"), summary.sample_count),
        (i18n.t("First sample_index"), "" if summary.first_sample_index is None else summary.first_sample_index),
        (i18n.t("Last sample_index"), "" if summary.last_sample_index is None else summary.last_sample_index),
        (i18n.t("Index span"), summary.index_span),
        (i18n.t("Samples missing (index span - stored)"), summary.missing_samples),
        (i18n.t("Heart-rate reports valid"), summary.hr_valid_reports),
        (i18n.t("Heart rate mean (valid)"), "" if summary.hr_mean is None else round(summary.hr_mean, 1)),
        (i18n.t("Heart rate min (valid)"), "" if summary.hr_min is None else summary.hr_min),
        (i18n.t("Heart rate max (valid)"), "" if summary.hr_max is None else summary.hr_max),
        (i18n.t("Temperature reports valid"), summary.temp_valid_reports),
        (i18n.t("Temperature mean (degC)"), "" if summary.temp_mean_c is None else round(summary.temp_mean_c, 2)),
        (i18n.t("Temperature min (degC)"), "" if summary.temp_min_c is None else round(summary.temp_min_c, 2)),
        (i18n.t("Temperature max (degC)"), "" if summary.temp_max_c is None else round(summary.temp_max_c, 2)),
        (
            i18n.t("Temperature certified?"),
            i18n.t("no (TEMP_UNCALIBRATED)") if summary.temp_uncalibrated else i18n.t("yes"),
        ),
        (
            i18n.t("Lead state"),
            i18n.t("UNKNOWN only (no lead-off hardware)")
            if summary.lead_unknown_only
            else i18n.t("see Data sheet"),
        ),
        (i18n.t("Probe state"), _probe_state_value(summary.probe_state_text)),
        (i18n.t("ECG raw min / max (ADC code)"), _both(summary.ecg_min_raw, summary.ecg_max_raw)),
        (i18n.t("Packets received"), summary.packets_total),
        (i18n.t("Sequence gaps"), summary.sequence_gaps),
        (i18n.t("Packets missing by sequence"), summary.sequence_missing),
        (i18n.t("ECG index gaps (batches lost)"), summary.index_gaps),
        (i18n.t("CRC errors"), summary.crc_errors),
        (i18n.t("Discarded bytes (resync)"), summary.discarded_bytes),
        (i18n.t("Protocol version"), "0x%02X" % protocol.PROTOCOL_VERSION),
        (
            i18n.t("Frame layout"),
            i18n.t("ECG_BATCH payload = %d + 2n bytes (tail = %d)")
            % (protocol.ECGP_SAMPLES, protocol.ECG_BATCH_TAIL_SIZE),
        ),
        (i18n.t("Exported at (PC clock)"), _iso(dt.datetime.now())),
        # The exporter is a package name plus a version: not a sentence.
        (i18n.t("Exporter"), "pc_monitor %s" % _tool_version()),
    ]


def export_xlsx(
    session: RecordingSession,
    path: str | Path,
    *,
    summary: RecordingSummary | None = None,
    max_chart_points: int = DEFAULT_CHART_POINTS,
) -> ExportResult:
    """Write the three-sheet workbook, including the embedded ECG line chart."""
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    summary = summary or session.summary()
    rows_written = 0

    wb = Workbook(write_only=True)

    # ---- Data: every row, nothing dropped -------------------------------
    data = wb.create_sheet(i18n.t(DATA_SHEET))
    data.freeze_panes = "A2"
    _set_widths(data, ROW_COLUMNS)
    data.append(_header_row(data, ROW_COLUMNS))
    for row in session.iter_rows():
        data.append(list(row))
        rows_written += 1

    # ---- ECG Series: decimated, chart-sized -----------------------------
    times, values = session.ecg_time_mv()
    points = recommended_chart_points(rows_written)
    dec_t, dec_v, per_point = decimate_for_chart(times, values, max_chart_points)
    series = wb.create_sheet(i18n.t(CHART_SHEET))
    # Translated even though nothing writes it today: the series sheet has to stay
    # the plain numeric range the chart points at, in either language.
    note = i18n.t(
        "Decimated for the chart only. The complete record is on '%s'. "
        "1 plotted point ~= %d sample(s); each bin contributes its maximum and "
        "minimum so the QRS survives."
    ) % (i18n.t(DATA_SHEET), max(1, per_point))
    series.append(_header_row(series, ("device_time_s", "ecg_pin_mv")))
    for t, v in zip(dec_t.tolist(), dec_v.tolist()):
        series.append([round(float(t), 6), round(float(v), 4)])
    _set_widths(series, ("device_time_s", "ecg_pin_mv"))

    # ---- Summary: statistics plus the chart -----------------------------
    info = wb.create_sheet(i18n.t(SUMMARY_SHEET))
    info.append([])
    title = WriteOnlyCell(info, value=i18n.t("ECG + body temperature recording -- summary"))
    title.font = Font(bold=True, size=14)
    info.append([title])
    if session.is_demo:
        warn = WriteOnlyCell(
            info, value=i18n.t("DEMO / SYNTHETIC DATA -- GENERATED ON THIS PC, NOT MEASURED")
        )
        warn.font = _WARNING_FONT
        warn.fill = _WARNING_FILL
        info.append([warn, warn, warn])
    info.append([])
    head_k = WriteOnlyCell(info, value=i18n.t("Item"))
    head_v = WriteOnlyCell(info, value=i18n.t("Value"))
    head_k.font = _HEADER_FONT
    head_k.fill = _HEADER_FILL
    head_v.font = _HEADER_FONT
    head_v.fill = _HEADER_FILL
    info.append([head_k, head_v])
    for key, value in _summary_pairs(summary, session):
        label = WriteOnlyCell(info, value=key)
        label.font = _LABEL_FONT
        info.append([label, value])

    info.append([])
    caveat_head = WriteOnlyCell(
        info, value=i18n.t("How to read this file (limits stated by the firmware contract)")
    )
    caveat_head.font = _LABEL_FONT
    info.append([caveat_head])
    for text in summary.caveats:
        cell = WriteOnlyCell(info, value=_caveat_text(text))
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        info.append([cell])
    info.column_dimensions["A"].width = 40
    info.column_dimensions["B"].width = 46

    chart = _build_chart(series, len(dec_t), session, per_point)
    # Anchored to the right of the text block, so it is the first thing seen.
    info.add_chart(chart, "D3")

    wb.save(str(target))
    return ExportResult(
        path=target,
        rows=rows_written,
        sheets=(DATA_SHEET, CHART_SHEET, SUMMARY_SHEET),
        chart_points=int(dec_v.size),
        chart_series_cells=int(dec_v.size),
    )


def _build_chart(series, point_count: int, session: RecordingSession, per_point: int) -> LineChart:
    """A real Excel line chart object bound to the ``ECG Series`` columns."""
    chart = LineChart()
    chart.title = i18n.t("%s ECG -- raw, millivolts at the MCU pin") % (
        i18n.t("DEMO / SYNTHETIC") if session.is_demo else i18n.t("Device")
    )
    chart.style = 2
    chart.height = 10
    chart.width = 34
    chart.x_axis.title = i18n.t("device time (s, from first_sample_index at 1 kHz)")
    chart.y_axis.title = i18n.t("ECG pin voltage (mV, unfiltered)")
    chart.x_axis.numFmt = "0.0"
    chart.y_axis.numFmt = "0"
    last = max(2, point_count + 1)
    values = Reference(series, min_col=2, min_row=1, max_row=last)
    categories = Reference(series, min_col=1, min_row=2, max_row=last)
    chart.add_data(values, titles_from_data=True)
    chart.set_categories(categories)
    plot = chart.series[0]
    plot.smooth = False  # Excel smoothing would lie about QRS width
    plot.marker.symbol = "none"
    if hasattr(plot, "tx"):
        # "DEMO" is a marker, not copy: it stays uppercase in either language.
        plot.tx = SeriesLabel(
            v=i18n.t("%s ECG mV") % ("DEMO" if session.is_demo else i18n.t("device"))
        )
    return chart


def _header_row(sheet, columns: Sequence[str]) -> list:
    """The styled header cells.  ``columns`` stay the field names: they are the
    keys into :data:`_WIDTHS` and the vocabulary of the wire protocol, so the
    translation is additive -- the sheet reads Chinese and still names the
    field."""
    cells = []
    for name in columns:
        cell = WriteOnlyCell(sheet, value=i18n.t(name))
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center")
        cells.append(cell)
    return cells


_WIDTHS = {
    "host_timestamp": 24,
    "host_epoch_s": 18,
    "device_ts_ms": 14,
    "sample_index": 13,
    "sample_time_s": 13,
    "ecg_raw": 9,
    "ecg_pin_mv": 12,
    "temp_raw": 10,
    "temp_centi": 11,
    "temp_c": 9,
    "hr_bpm": 8,
    "hr_valid": 9,
    "hr_state": 9,
    "lead_state": 15,
    "probe_state": 17,
    "status_flags": 12,
    "origin": 20,
    "device_time_s": 14,
}


def _set_widths(sheet, columns: Iterable[str]) -> None:
    for offset, name in enumerate(columns, start=1):
        sheet.column_dimensions[get_column_letter(offset)].width = _WIDTHS.get(name, 14)


def _iso(moment: dt.datetime | None) -> str:
    return moment.isoformat(timespec="seconds") if moment else i18n.t("n/a")


def _both(low, high) -> str:
    if low is None or high is None:
        return i18n.t("n/a")
    return "%s / %s" % (low, high)


#: The two phrases :func:`pc_monitor.recorder._probe_text` adds around the
#: ``TempState`` names.  The enum names themselves stay English in either
#: language, so only these are looked up, and the key carries the space so the
#: English cell stays byte-identical to what the recorder builds.
_PROBE_NO_DATA = "no data"
_PROBE_UNCERTIFIED = " (temperature uncertified)"


def _probe_state_value(text: str) -> str:
    """Translate the host-written part of a probe-state string, names included."""
    if text == _PROBE_NO_DATA:
        return i18n.t(_PROBE_NO_DATA)
    if text.endswith(_PROBE_UNCERTIFIED):
        return text[: -len(_PROBE_UNCERTIFIED)] + i18n.t(_PROBE_UNCERTIFIED)
    return text


#: ``RecordingSession.caveats()`` hands back finished sentences, one of which has
#: a count already interpolated -- so that note is matched, not looked up, and
#: translated from its template.  A wording change on the recorder side falls
#: through to :func:`pc_monitor.i18n.t` verbatim, which records it as missed.
_HR_DISAGREEMENT_NOTE = (
    "%d batch(es) disagreed about heart-rate validity across hr_bpm / hr_state / "
    "SFLAG_HR_VALID; those were recorded as invalid."
)
_COUNT_PREFIXED_NOTE = re.compile(r"^(\d+) (batch\(es\) disagreed about )")


def _caveat_text(text: str) -> str:
    """One line of the 'how to read this file' block, in the active language."""
    match = _COUNT_PREFIXED_NOTE.match(text)
    if match:
        return i18n.t(_HR_DISAGREEMENT_NOTE) % int(match.group(1))
    return i18n.t(text)


def _capabilities(session: RecordingSession) -> str:
    hello = getattr(session, "hello", None)
    if hello is None:
        return i18n.t("no HELLO received: every capability is treated as absent")
    caps = protocol.Capability(hello.caps)
    named = [cap.name for cap in protocol.Capability if cap is not protocol.Capability.NONE and caps & cap]
    missing = [
        cap.name
        for cap in protocol.Capability
        if cap is not protocol.Capability.NONE and not caps & cap
    ]
    # ``Capability`` member names are protocol vocabulary: they stay English here
    # so the line can be read straight against protocol.h.
    return i18n.t("present: %s | absent: %s") % (
        ", ".join(named) or i18n.t("none"),
        ", ".join(missing) or i18n.t("none"),
    )


def _tool_version() -> str:
    from . import __version__

    return __version__


def readback_workbook(path: str | Path):
    """Open a written workbook -- used by the tests to prove the file is valid."""
    return load_workbook(str(path))
